"""Comprehensive test suite for sciTree core, CLI, HTTP API, MCP tools, extraction, and export."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import docx
try:
    import pymupdf as fitz
except ImportError:
    import fitz
import pytest
from fastapi.testclient import TestClient
from typer.testing import CliRunner

from scitree.api.server import create_app
from scitree.cli import app as cli_app
from scitree.config import BootstrapStateManager, get_recommended_root_path
from scitree.core import SciTreeCore
from scitree.errors import (
    ERR_ACTOR_REQUIRED,
    ERR_ROOT_NOT_CONFIGURED,
    ERR_TEMPLATE_NOT_FOUND,
    ERR_UNAUTHORIZED,
    ERR_VERSION_CONFLICT,
    ERR_VERSION_IMMUTABLE,
    SciTreeError,
)
from scitree.mcp_server import (
    add_comment as mcp_add_comment,
    add_milestone as mcp_add_milestone,
    add_risk as mcp_add_risk,
    add_version as mcp_add_version,
    apply_manifest as mcp_apply_manifest,
    create_project as mcp_create_project,
    export_package as mcp_export_package,
    init_root as mcp_init_root,
    list_projects as mcp_list_projects,
)
from scitree.models import validate_actor

PNG_1X1_BYTES = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
    b"\x00\x00\x03\x01\x01\x00\x18\xdd\x8d\xb0\x00\x00\x00\x00IEND\xaeB`\x82"
)


@pytest.fixture()
def workspace_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Path]:
    state_file = tmp_path / "config" / "bootstrap.json"
    root_dir = tmp_path / "scitree_root"
    assets_dir = tmp_path / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("SCITREE_STATE_FILE", str(state_file))

    # Create sample PNG, CSV, DOCX (with image + inline comment + word/comments.xml + references), and PDF
    img_path = assets_dir / "fig1.png"
    img_path.write_bytes(PNG_1X1_BYTES)

    csv_path = assets_dir / "dataset.csv"
    csv_path.write_text("id,val\n1,42.0\n2,43.5\n", encoding="utf-8")

    docx_v1 = assets_dir / "manuscript_v1.docx"
    doc = docx.Document()
    doc.core_properties.author = "Jane Doe"
    doc.add_heading("Manuscript Title", level=1)
    doc.add_paragraph("Body paragraph with [COMMENT: Verify statistical model assumptions.] inline.")
    doc.add_picture(str(img_path))
    doc.add_heading("References", level=1)
    doc.add_paragraph("Doe, J., & Smith, J. (2024). Neural dynamics of sleep. Sleep Science, 12(1), 10-22.")
    doc.add_paragraph("Silva, M. (2023). Circadian rhythms in students. Chronobiology International, 40(4), 300-312.")
    doc.save(str(docx_v1))

    # Inject a native word/comments.xml into docx_v1
    comments_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:comment w:id="0" w:author="John Smith"><w:p><w:r><w:t>Clarify inclusion criteria in methods.</w:t></w:r></w:p></w:comment>'
        "</w:comments>"
    )
    with zipfile.ZipFile(docx_v1, "a") as zf:
        zf.writestr("word/comments.xml", comments_xml.encode("utf-8"))

    # Create a second distinct DOCX for conflict testing
    docx_v2 = assets_dir / "manuscript_v2.docx"
    doc2 = docx.Document()
    doc2.add_heading("Revised Manuscript Title", level=1)
    doc2.add_paragraph("Updated body text after peer review.")
    doc2.add_heading("References", level=1)
    doc2.add_paragraph("Doe, J. (2025). Follow-up sleep study. Sleep, 48(1), 1-9.")
    doc2.save(str(docx_v2))

    # Create a sample PDF with image, annotation, inline marker, and references
    pdf_path = assets_dir / "manuscript.pdf"
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((72, 72), "PDF Manuscript Title\n[COMMENT(agent:claude): Check Figure 2 axis labels.]\nReferences\n1. Brown, A. (2022). PDF extraction workflows. J. Comp. Sci., 5, 12-19.")
    page.insert_image(fitz.Rect(72, 140, 120, 188), stream=PNG_1X1_BYTES)
    annot = page.add_text_annot(fitz.Point(72, 100), "PDF reviewer annotation comment.")
    annot.set_info(title="Maria Silva", content="PDF reviewer annotation comment.")
    annot.update()
    pdf_doc.save(str(pdf_path))
    pdf_doc.close()

    return {
        "state_file": state_file,
        "root": root_dir,
        "assets": assets_dir,
        "img": img_path,
        "csv": csv_path,
        "docx_v1": docx_v1,
        "docx_v2": docx_v2,
        "pdf": pdf_path,
    }


def test_root_onboarding_and_os_defaults(workspace_env: dict[str, Path]) -> None:
    win_path = get_recommended_root_path("Windows")
    assert win_path.name == "sciTree" and "Documents" in str(win_path)

    mac_path = get_recommended_root_path("Darwin")
    assert mac_path.name == "sciTree" and "Documents" in str(mac_path)

    linux_path = get_recommended_root_path("Linux")
    assert linux_path.name.lower() == "scitree"

    core = SciTreeCore()
    with pytest.raises(SciTreeError) as exc_info:
        core.list_projects()
    assert exc_info.value.code == ERR_ROOT_NOT_CONFIGURED

    # Initialize root
    init_res = core.init_root(root_path=workspace_env["root"])
    assert init_res["initialized"] is True
    assert (workspace_env["root"] / "Projects").is_dir()
    assert (workspace_env["root"] / "Templates" / "Empirical_Paper.json").is_file()
    assert (workspace_env["root"] / "Automation" / "inbox").is_dir()
    assert (workspace_env["root"] / "Exports").is_dir()

    # Reset onboarding
    reset_res = core.reset_onboarding()
    assert reset_res["onboarding_completed"] is False


def test_project_templates_versions_immutability_and_extraction(workspace_env: dict[str, Path]) -> None:
    core = SciTreeCore(root_path=workspace_env["root"])
    core.init_root(root_path=workspace_env["root"])

    # Idempotent project creation with Empirical Paper template
    res1 = core.create_project(name="Alpha Study", template_name="Empirical Paper", column="Writing")
    assert res1["created"] is True
    assert "Literature" in res1["project"]["status"]["column_order"]
    assert len(res1["project"]["milestones"]) >= 4

    res2 = core.create_project(name="Alpha Study", template_name="Empirical Paper")
    assert res2["created"] is False
    assert res2["idempotent"] is True

    # Upload DOCX version v1
    v1_res = core.add_version(
        project_name="Alpha Study",
        file_path=workspace_env["docx_v1"],
        author="Jane Doe",
        version_id="v1",
        changelog="Initial draft",
    )
    assert v1_res["created"] is True
    assert v1_res["version"]["extraction"]["figures_extracted"] == 1
    assert v1_res["version"]["extraction"]["references_parsed"] == 2
    assert len(v1_res["project"]["comments"]) == 2  # 1 docx_comment + 1 inline_marker

    # Re-submitting identical v1 is idempotent
    v1_idem = core.add_version(
        project_name="Alpha Study",
        file_path=workspace_env["docx_v1"],
        author="Jane Doe",
        version_id="v1",
    )
    assert v1_idem["created"] is False
    assert v1_idem["idempotent"] is True

    # Submitting different file with same version_id v1 raises VERSION_CONFLICT with suggested suffix v1a
    with pytest.raises(SciTreeError) as conflict_exc:
        core.add_version(
            project_name="Alpha Study",
            file_path=workspace_env["docx_v2"],
            author="Jane Doe",
            version_id="v1",
        )
    assert conflict_exc.value.code == ERR_VERSION_CONFLICT
    assert conflict_exc.value.details["suggested_version_id"] == "v1a"

    # Upload PDF as v2
    v2_res = core.add_version(
        project_name="Alpha Study",
        file_path=workspace_env["pdf"],
        author="Maria Silva",
        version_id="v2",
        changelog="PDF submission draft",
    )
    assert v2_res["created"] is True
    assert v2_res["version"]["extraction"]["figures_extracted"] == 1
    assert v2_res["version"]["extraction"]["references_parsed"] == 1

    # Promote v2 to submitted -> becomes immutable
    lock_res = core.update_version_state(
        project_name="Alpha Study",
        version_id="v2",
        state="submitted",
        submitted_to="Journal of Sleep Research",
    )
    assert lock_res["version"]["immutable"] is True
    assert lock_res["project"]["latest_version_locked"] is True

    # Attempting to overwrite or modify immutable v2 raises VERSION_IMMUTABLE
    with pytest.raises(SciTreeError) as imm_exc:
        core.add_version(
            project_name="Alpha Study",
            file_path=workspace_env["docx_v2"],
            author="Jane Doe",
            version_id="v2",
        )
    assert imm_exc.value.code == ERR_VERSION_IMMUTABLE
    assert imm_exc.value.details["derived_from"] == "v2"

    # Deleting immutable v2 is forbidden
    with pytest.raises(SciTreeError) as del_exc:
        core.delete_version(project_name="Alpha Study", version_id="v2")
    assert del_exc.value.code == ERR_VERSION_IMMUTABLE

    # Creating subsequent v3 automatically links derived_from="v2"
    v3_res = core.add_version(
        project_name="Alpha Study",
        file_path=workspace_env["docx_v2"],
        author="Jane Doe",
    )
    assert v3_res["version"]["version_id"] == "v3"
    assert v3_res["version"]["derived_from"] == "v2"


def test_risks_milestones_export_reproducibility_and_db_rebuild(workspace_env: dict[str, Path]) -> None:
    core = SciTreeCore(root_path=workspace_env["root"])
    core.init_root(root_path=workspace_env["root"])
    core.create_project(name="Beta Trial")
    core.add_version(
        project_name="Beta Trial",
        file_path=workspace_env["docx_v1"],
        author="Jane Doe",
        version_id="v1",
    )

    # Add risk & close risk
    r_res = core.add_risk(
        project_name="Beta Trial",
        title="Equipment calibration drift",
        severity="critical",
        likelihood="medium",
        risk_id="rsk_custom",
    )
    assert r_res["created"] is True
    assert r_res["project"]["risks_summary"]["high_or_critical_open"] == 1

    # Idempotent risk add
    r_idem = core.add_risk(
        project_name="Beta Trial",
        title="Equipment calibration drift",
        risk_id="rsk_custom",
    )
    assert r_idem["idempotent"] is True

    core.close_risk(project_name="Beta Trial", risk_id="rsk_custom")
    assert core.get_project("Beta Trial")["project"]["risks_summary"]["high_or_critical_open"] == 0

    # Add milestone & mark done
    ms_res = core.add_milestone(
        project_name="Beta Trial",
        title="Ethics approval",
        phase="short_term_7d",
        due_date="2026-11-01",
        milestone_id="ms_custom",
    )
    assert ms_res["created"] is True
    core.close_milestone(project_name="Beta Trial", milestone_id="ms_custom")
    assert core.get_project("Beta Trial")["project"]["milestones_summary"]["done"] >= 1

    # Reproducible submission package export (Section F2)
    exp1 = core.export_package(
        project_name="Beta Trial",
        version_id="v1",
        output_path=workspace_env["root"] / "Exports" / "pkg1.zip",
    )
    exp2 = core.export_package(
        project_name="Beta Trial",
        version_id="v1",
        output_path=workspace_env["root"] / "Exports" / "pkg2.zip",
    )
    assert exp1["export"]["checksum"] == exp2["export"]["checksum"]

    # Rebuild SQLite database from filesystem (Section I2)
    stats = core._get_db(workspace_env["root"]).rebuild_from_filesystem(workspace_env["root"])
    assert stats["projects_indexed"] == 1
    assert stats["templates_indexed"] >= 4

    # Actor attribution requirement (Section I5)
    with pytest.raises(SciTreeError) as actor_exc:
        validate_actor(actor_name="", actor_type="agent", actor_source="cli", require_agent_name=True)
    assert actor_exc.value.code == ERR_ACTOR_REQUIRED


def test_cli_http_mcp_manifest_and_inbox(workspace_env: dict[str, Path]) -> None:
    runner = CliRunner()
    root_str = str(workspace_env["root"])

    # 1. CLI tests
    res_init = runner.invoke(cli_app, ["init", "--root", root_str, "--json"])
    assert res_init.exit_code == 0
    assert json.loads(res_init.stdout)["initialized"] is True

    res_create = runner.invoke(
        cli_app,
        ["project", "create", "--name", "CLI Study", "--template", "Review Article", "--json"],
    )
    assert res_create.exit_code == 0
    assert json.loads(res_create.stdout)["project"]["name"] == "CLI Study"

    # Dry-run check
    res_dry = runner.invoke(
        cli_app,
        ["project", "create", "--name", "DryRunProject", "--dry-run", "--json"],
    )
    assert res_dry.exit_code == 0
    assert json.loads(res_dry.stdout)["dry_run"] is True
    res_list = runner.invoke(cli_app, ["project", "list", "--json"])
    names = [p["name"] for p in json.loads(res_list.stdout)["projects"]]
    assert "DryRunProject" not in names

    # 2. HTTP API tests
    api_app = create_app(default_root=root_str)
    client = TestClient(api_app)

    r_proj = client.post("/projects", json={"name": "HTTP Study", "column": "Idea", "template": "Grant Proposal"})
    assert r_proj.status_code == 200
    assert r_proj.json()["project"]["name"] == "HTTP Study"

    r_ver = client.post(
        "/projects/HTTP Study/versions",
        json={"file": str(workspace_env["docx_v1"]), "author": "Dr. HTTP", "version_id": "v1"},
    )
    assert r_ver.status_code == 200
    assert r_ver.json()["version"]["version_id"] == "v1"

    r_exp = client.post("/projects/HTTP Study/export", json={"version": "v1"})
    assert r_exp.status_code == 200
    assert "checksum" in r_exp.json()["export"]

    # 3. MCP tool tests
    mcp_res = mcp_create_project(name="MCP Study", root=root_str, actor_name="agent:mcp-test")
    assert mcp_res["project"]["name"] == "MCP Study"

    mcp_v = mcp_add_version(
        project="MCP Study",
        file=str(workspace_env["pdf"]),
        author="agent:mcp-test",
        version_id="v1",
        root=root_str,
    )
    assert mcp_v["version"]["version_id"] == "v1"
    mcp_c = mcp_add_comment(project="MCP Study", text="MCP review comment", root=root_str)
    assert mcp_c["comment"]["owner_type"] == "agent"

    # 4. Watch folder (Automation/inbox -> Automation/outbox)
    inbox_file = workspace_env["root"] / "Automation" / "inbox" / "job1.json"
    inbox_file.write_text(
        json.dumps(
            {
                "project": {"name": "Inbox Study", "column": "Writing"},
                "risks": [{"title": "Inbox risk item", "severity": "high"}],
                "milestones": [{"title": "Inbox milestone", "phase": "short_term_7d"}],
            }
        ),
        encoding="utf-8",
    )
    r_inbox = client.post("/inbox/process")
    assert r_inbox.status_code == 200
    assert r_inbox.json()["processed_count"] == 1
    outbox_result = workspace_env["root"] / "Automation" / "outbox" / "job1.json.result.json"
    assert outbox_result.is_file()
